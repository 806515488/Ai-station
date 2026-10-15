"""附件真实离线回归：临时文件/SQLite，子进程禁止联网，不调云模型。"""
from __future__ import annotations

import io
import json
import time
import zipfile

import pytest
from fastapi.testclient import TestClient

from station import db
from station.app import server
from station.documents import service, settings, storage
from station.documents.cleaning import clean
from station.jobs.manager import get_manager


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("STATION_FAKE", "1")
    monkeypatch.setattr(server, "l2_route", lambda *_args, **_kw: None)
    with TestClient(server.app) as client:
        yield client
    # 夹具切换数据目录前必须让后台任务退出，不能把尾部结果写进下一个测试库。
    deadline = time.monotonic() + 30
    while get_manager()._active and time.monotonic() < deadline:
        time.sleep(0.05)
    assert not get_manager()._active


def login(client, name="documents-a"):
    response = client.post("/api/auth/register", json={"name": name, "password": "p"})
    assert response.status_code == 200, response.text
    return response.json()["user"]["id"]


def upload(client, name, content, key="upload-a"):
    response = client.post("/api/documents/upload", data={"client_upload_id": key},
                           files=[("files", (name, content))])
    assert response.status_code == 200, response.text
    data = response.json()
    assert not data["errors"], data
    return data["documents"][0]


def events(response):
    assert response.status_code == 200, response.text
    return [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: ")]


def start(client, documents, message="请总结附件", key="turn-a", thread_id=None):
    rows = events(client.post("/api/chat", json={"message": message,
                  "attachment_ids": [d["id"] for d in documents], "client_turn_id": key,
                  "thread_id": thread_id}))
    assert not [r for r in rows if r["type"] == "error"], rows
    return next(r for r in rows if r["type"] == "attachment_wait")


def finish(client, turn_id):
    deadline = time.monotonic() + 90
    while time.monotonic() < deadline:
        response = client.get("/api/documents/turns/" + turn_id)
        assert response.status_code == 200
        value = response.json()
        if value["job"]["status"] in {"done", "failed"}:
            return value
        time.sleep(0.1)
    raise AssertionError("附件任务超时")


def manifest(owner, document_id):
    return storage.manifest(db.document_get(document_id, owner))


def test_plaintext_round_trip_cleaning_chunks_and_model_inputs(client, monkeypatch):
    owner = login(client)
    raw = "报销制度：\n\n金额 -120 元，不能省略否定词。\n\n" * 100
    document = upload(client, "制度.txt", raw.encode())
    turn = start(client, [document])
    done = finish(client, turn["turn_id"])
    assert done["state"] == "ready" and not done["failures"], done
    folder, value = manifest(owner, document["id"])
    blocks = list(storage.records(folder, value["parsed"]["shards"]))
    assert "".join(b["text"] for b in blocks) == raw
    chunks = list(storage.records(folder, value["chunk_sets"][0]["shards"]))
    assert len(chunks) > 1 and all(len(c["text"]) <= 1000 for c in chunks)
    by_id = {b["block_id"]: b for b in blocks}
    for chunk in chunks:
        assert "".join(by_id[s["block_id"]]["text"][s["start"]:s["end"]] for s in chunk["source_spans"]) == chunk["text"]
    captured = []
    class Model:
        def stream(self, messages, _tools):
            captured.extend(messages)
            yield {"final": {"content": "已阅读", "tool_calls": []}}
    monkeypatch.setattr(server, "_make_model", lambda *_a: Model())
    rows = events(client.post("/api/chat", json={"message": "", "thread_id": turn["thread_id"],
                                                "resume_turn_id": turn["turn_id"]}))
    assert any(r["type"] == "done" for r in rows)
    assert any(raw[:100] in m.get("content", "") for m in captured)
    snapshot = client.get("/api/sessions/" + turn["thread_id"]).json()
    assert [m["content"] for m in snapshot["msgs"] if m["role"] == "user"] == ["请总结附件"]
    # 重复恢复不能多跑一次模型，也不能重复追加用户问题。
    captured.clear()
    client.post("/api/chat", json={"message": "", "resume_turn_id": turn["turn_id"]})
    assert not captured
    assert client.get("/api/files").json()["documents"][0]["date"] == db.document_get(document["id"], owner)["date"]


def test_same_bytes_two_users_are_isolated(client):
    owner_a = login(client)
    a = upload(client, "相同.txt", b"same source")
    done_a = start(client, [a])
    finish(client, done_a["turn_id"])
    details = client.get("/api/documents/" + a["id"]).json()
    resource_url = details["files"][0]["url"]
    client.cookies.clear()
    owner_b = login(client, "documents-b")
    b = upload(client, "相同.txt", b"same source")
    assert a["id"] != b["id"] and owner_a != owner_b
    assert client.get(resource_url).status_code == 404
    assert client.get("/api/documents/turns/" + done_a["turn_id"]).status_code == 404
    assert client.get("/api/documents/" + a["id"]).status_code == 404
    assert not any(x["id"] == a["id"] for x in client.get("/api/files").json()["documents"])


def test_upload_replay_and_fixed_turn(client):
    login(client)
    a = upload(client, "a.txt", b"one")
    again = upload(client, "a.txt", b"one")
    assert a["id"] == again["id"]
    turn = start(client, [a])
    # 首次响应丢失时，没有 thread_id 的重发也必须命中原轮次。
    second = start(client, [a])
    assert second["turn_id"] == turn["turn_id"]
    finish(client, turn["turn_id"])
    bad = events(client.post("/api/chat", json={"message": "different", "attachment_ids": [a["id"]],
                         "client_turn_id": "turn-a", "thread_id": turn["thread_id"]}))
    assert any(r["type"] == "error" for r in bad)


def test_markdown_code_and_chinese_preserved():
    value = "# 第一章\n\n```python\nif x < 0:\n    print('中文 -12')\n```\n\n|值|单位|\n|-|-|\n|1|元|"
    result, transforms = clean(value, ".md", {})
    assert result == value and transforms == []
    paragraph, _ = clean("中文    -12 元。", ".html", {"category": "NarrativeText"})
    assert paragraph == "中文 -12 元。"


def test_long_document_shards_cache_and_ids_stable(client, monkeypatch):
    owner = login(client)
    raw = ("第1条：本次费用 -120 元，不能重复报销。\n\n" * 6000).encode()
    first = upload(client, "长文.txt", raw)
    monkeypatch.setenv("STATION_DOCUMENT_SHARD_BYTES", "32768")
    t = start(client, [first])
    finish(client, t["turn_id"])
    folder, value = manifest(owner, first["id"])
    assert len(value["parsed"]["shards"]) > 1
    ids = [c["chunk_id"] for c in storage.records(folder, value["chunk_sets"][0]["shards"])]
    monkeypatch.setenv("STATION_DOCUMENT_SHARD_BYTES", "1048576")
    retry = client.post("/api/documents/" + first["id"] + "/process")
    assert retry.status_code == 200, retry.text
    finish(client, retry.json()["id"])
    folder2, value2 = manifest(owner, first["id"])
    assert value2["processing"]["cache_hit"] is True
    assert [c["chunk_id"] for c in storage.records(folder2, value2["chunk_sets"][0]["shards"])] == ids
    assert value["parsed"]["snapshot_id"] == value2["parsed"]["snapshot_id"]
    package = client.get("/api/documents/" + first["id"] + "/archive")
    assert package.status_code == 200
    with zipfile.ZipFile(io.BytesIO(package.content)) as archive:
        assert "manifest.json" in archive.namelist()
        assert archive.read("original/长文.txt") == raw


def test_real_docx_csv_pdf_and_html_loaders_offline(client):
    from docx import Document
    from reportlab.pdfgen.canvas import Canvas
    from openpyxl import Workbook
    owner = login(client)
    word = io.BytesIO()
    doc = Document(); doc.add_paragraph("中文合同 -120 元，必须保留。")
    doc.save(word)
    pdf = io.BytesIO()
    canvas = Canvas(pdf); canvas.drawString(30, 700, "Page one: -120 USD"); canvas.showPage()
    canvas.drawString(30, 700, "Page two: keep units"); canvas.save()
    sheet = io.BytesIO(); workbook = Workbook()
    workbook.active.append(["地区", "金额（元）"])
    workbook.active.append(["华东", -120]); workbook.save(sheet)
    docs = [upload(client, "合同.docx", word.getvalue(), "word"),
            upload(client, "报表.csv", "地区,金额\n华东,-120\n华南,30\n".encode(), "csv"),
            upload(client, "manual.pdf", pdf.getvalue(), "pdf"),
            upload(client, "page.html", "<html><body><p>中文资料需要保留 -120 元。</p></body></html>".encode(), "html"),
            upload(client, "表格.xlsx", sheet.getvalue(), "xlsx")]
    t = start(client, docs)
    done = finish(client, t["turn_id"])
    assert not done["failures"], done
    for document in docs:
        folder, value = manifest(owner, document["id"])
        blocks = list(storage.records(folder, value["parsed"]["shards"]))
        assert any("-120" in b["text"] for b in blocks), document
        if document["suffix"] == ".pdf":
            assert [b["source_locations"][0]["page_index"] for b in blocks] == [1, 2]


def test_scan_pdf_is_not_false_success(client):
    from PIL import Image
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas
    owner = login(client)
    pdf = io.BytesIO(); canvas = Canvas(pdf)
    canvas.drawImage(ImageReader(Image.new("RGB", (30, 30), "black")), 30, 600)
    canvas.save()
    document = upload(client, "扫描.pdf", pdf.getvalue())
    turn = start(client, [document]); finish(client, turn["turn_id"])
    folder, value = manifest(owner, document["id"])
    assert value["quality"]["status"] == "unusable"
    assert any(x["code"] == "needs_ocr" for x in value["quality"]["issues"])
    assert db.document_get(document["id"], owner)["status"] == "failed"


def test_mixed_failure_does_not_destroy_good_result(client):
    owner = login(client)
    good = upload(client, "good.txt", b"useful content", "good")
    broken = upload(client, "broken.docx", b"PK malformed zip", "bad")
    turn = start(client, [good, broken]); done = finish(client, turn["turn_id"])
    assert broken["id"] in done["failures"] and good["id"] not in done["failures"]
    assert manifest(owner, good["id"])[1]["quality"]["status"] == "accepted"


def test_shard_checksum_rejects_tampering(client):
    owner = login(client)
    doc = upload(client, "a.txt", b"content")
    t = start(client, [doc]); finish(client, t["turn_id"])
    folder, value = manifest(owner, doc["id"])
    path = folder / value["parsed"]["shards"][0]["path"]
    record = json.loads(path.read_text(encoding="utf-8"))
    record["text"] = "tampered but valid JSON"
    path.write_text(json.dumps(record, ensure_ascii=False) + "\n", encoding="utf-8")
    with pytest.raises(ValueError, match="校验失败"):
        list(storage.records(folder, value["parsed"]["shards"]))


def test_resource_fingerprint_ignores_physical_shards():
    a = settings.snapshot(); b = settings.snapshot()
    b["limits"]["shard_bytes"] *= 2
    assert settings.fingerprints("source", a, {}) == settings.fingerprints("source", b, {})


def test_concurrent_resume_does_not_interrupt_the_winner(client, monkeypatch):
    import threading
    owner = login(client)
    doc = upload(client, "a.txt", b"some content")
    turn = start(client, [doc]); finish(client, turn["turn_id"])
    entered, release = threading.Event(), threading.Event()
    calls = []
    class SlowModel:
        def stream(self, _messages, _tools):
            calls.append(1); entered.set(); assert release.wait(10)
            yield {"final": {"content": "完成", "tool_calls": []}}
    monkeypatch.setattr(server, "_make_model", lambda *_args: SlowModel())
    responses = []
    payload = {"message": "", "resume_turn_id": turn["turn_id"]}
    worker = threading.Thread(target=lambda: responses.append(client.post("/api/chat", json=payload)))
    worker.start(); assert entered.wait(10)
    try:
        second = client.post("/api/chat", json=payload)
        assert second.status_code == 200
        assert db.attachment_turn_get(turn["turn_id"], owner)["state"] == "running_agent"
    finally:
        release.set(); worker.join(10)
    assert not worker.is_alive() and len(calls) == 1
    assert db.attachment_turn_get(turn["turn_id"], owner)["state"] == "done"


def test_failed_model_keeps_parsed_and_explicit_retry(client, monkeypatch):
    owner = login(client)
    doc = upload(client, "a.txt", b"facts")
    turn = start(client, [doc]); finish(client, turn["turn_id"])
    def unavailable(*_a):
        raise RuntimeError("missing model")
    monkeypatch.setattr(server, "_make_model", unavailable)
    payload = {"message": "", "resume_turn_id": turn["turn_id"]}
    assert any(e["type"] == "error" for e in events(client.post("/api/chat", json=payload)))
    assert db.attachment_turn_get(turn["turn_id"], owner)["state"] == "failed"
    assert manifest(owner, doc["id"])[1]["parsed"]["status"] == "ready"
    # 默认恢复不能悄悄重新执行模型；只有显式 retry_agent 才能重新领取。
    assert not any(e["type"] == "error" for e in events(client.post("/api/chat", json=payload)))
    class RetryModel:
        def stream(self, _messages, _tools):
            yield {"final": {"content": "恢复完成", "tool_calls": []}}
    monkeypatch.setattr(server, "_make_model", lambda *_a: RetryModel())
    payload["retry_agent"] = True
    assert any(e["type"] == "done" for e in events(client.post("/api/chat", json=payload)))
    assert db.attachment_turn_get(turn["turn_id"], owner)["state"] == "done"


def test_attachment_replay_position_survives_card_pruning():
    rows = server._replay_items([{"role": "user", "content": "问题", "attachment_turn_id": "turn"},
                                 {"role": "assistant", "content": "回答"}], [])
    assert [row["role"] for row in rows] == ["user", "render", "assistant"]
    assert rows[1]["card"]["turn_id"] == "turn"


def test_postgres_existing_schema_requires_document_upgrade():
    class Cursor:
        def __init__(self, value): self.value = value
        def fetchall(self): return self.value
        def fetchone(self): return self.value
    class Connection:
        def __init__(self, tables): self.tables = tables
        def execute(self, statement):
            if "information_schema.tables" in statement:
                return Cursor([{"table_name": name} for name in self.tables])
            return Cursor({"present": True})
    old = db._PG_CORE_TABLES | db._PG_INTERVIEW_TABLES
    assert db._postgres_schema_state(Connection(old)) == "upgrade-needed"
    assert db._postgres_schema_state(Connection(db._PG_TABLES)) == "current"


def test_cancel_queued_processing_keeps_original(client):
    from station.core.heavy import heavy_slot
    owner = login(client)
    doc = upload(client, "a.txt", b"original")
    with heavy_slot("test blocker"):
        turn = start(client, [doc])
        response = client.post("/api/documents/turns/" + turn["turn_id"] + "/cancel")
        assert response.status_code == 200
    done = finish(client, turn["turn_id"])
    assert done["state"] == "cancelled"
    original = client.get("/api/documents/" + doc["id"]).json()["files"][0]
    assert client.get(original["url"]).content == b"original"
    assert db.document_get(doc["id"], owner) is not None


def test_restart_recovers_persisted_parse_job(client, monkeypatch):
    from station.jobs.manager import Job
    owner = login(client)
    doc = upload(client, "a.txt", b"recover locally")
    manager = get_manager()
    def persist_only(executor, args, user_id, job_id=""):
        job = Job(executor.id, args); job.id = job_id; job.user_id = user_id
        job._persist()
        return job
    with monkeypatch.context() as patch:
        patch.setattr(manager, "submit", persist_only)
        turn = start(client, [doc])
    assert db.attachment_turn_get(turn["turn_id"], owner)["state"] == "waiting_parse"
    service.recover()
    done = finish(client, turn["turn_id"])
    assert done["state"] == "ready" and not done["failures"]
    messages = client.get("/api/sessions/" + turn["thread_id"]).json()["msgs"]
    assert sum(m["role"] == "user" for m in messages) == 1
    db.attachment_turn_update(turn["turn_id"], owner,
                              {"state": "running_agent", "agent_boot_id": "previous-process"})
    service.recover()
    assert db.attachment_turn_get(turn["turn_id"], owner)["state"] == "interrupted"


def test_only_unused_drafts_can_be_removed(client):
    owner = login(client)
    draft = upload(client, "unused.txt", b"unused")
    folder = storage.root(owner, draft["id"])
    assert client.delete("/api/documents/" + draft["id"]).status_code == 200
    assert db.document_get(draft["id"], owner) is None and not folder.exists()
    used = upload(client, "used.txt", b"facts", "used")
    turn = start(client, [used]); finish(client, turn["turn_id"])
    assert client.delete("/api/documents/" + used["id"]).status_code == 409
    details = client.get("/api/documents/" + used["id"]).json()
    part = next(f for f in details["files"] if f["name"].startswith("parsed/") and f["name"].endswith("jsonl"))
    preview = client.get(part["url"] + "?preview=1").json()
    assert "facts" in preview["text"] and not preview["truncated"]


def test_schema_rejects_chunks_without_sources(client):
    from station.documents.quality import validate
    owner = login(client)
    doc = upload(client, "a.txt", b"facts")
    turn = start(client, [doc]); finish(client, turn["turn_id"])
    folder, value = manifest(owner, doc["id"])
    chunk = next(storage.records(folder, value["chunk_sets"][0]["shards"]))
    validate("chunk", chunk)
    del chunk["source_spans"]
    with pytest.raises(ValueError, match="契约校验失败"):
        validate("chunk", chunk)


def test_empty_document_resource_option_uses_packaged_nltk(client, monkeypatch):
    import os
    resource_path = os.environ.get("STATION_DOCUMENT_NLTK_DATA")
    assert resource_path, "本地文档资源须先准备"
    monkeypatch.setenv("NLTK_DATA", resource_path)
    monkeypatch.setenv("STATION_DOCUMENT_NLTK_DATA", "")
    owner = login(client)
    doc = upload(client, "a.html", b"<html><body><p>Keep -120 USD and important facts.</p></body></html>")
    turn = start(client, [doc]); done = finish(client, turn["turn_id"])
    assert not done["failures"], done
    folder, value = manifest(owner, doc["id"])
    assert any("-120" in b["text"] for b in storage.records(folder, value["parsed"]["shards"]))
