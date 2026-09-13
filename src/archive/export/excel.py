"""⑤ Excel 导出：严格复刻《赵景刚.xlsx》模板。

策略：先输出与模板完全一致的固定骨架（十个大类、四/九下的小类行、
表头、边框、字体、行高、列宽、尾部空行），再把识别出的材料行
填入对应大类/小类之下。无材料的类别保留骨架行（与模板一致）。

新手视角（Java 朋友版）：excel 导出 ≈ 用 Apache POI 生成固定样式 Excel 的报表类。
  思路 = “先画模板骨架（表头/十类行/四九小类行/边框/字体）→ 再把数据填进对应行”。
  你只要看懂两个点：① SKELETON（骨架表，见下方常量）② _write_materials 里的
  _date_key（材料按制成时间排序的规则）。
  为什么“手工造表”而不是用模板文件？为了像素级一致 + 可控边框字体。
"""
from __future__ import annotations

import os                            # 建目录/拼路径用

# openpyxl = Python 操作 .xlsx 的库。Workbook≈一个 Excel 文件，Sheet≈一张表。
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Border, Side   # 样式对象：对齐/字体/边框/边

from archive.domain import classes          # 目录序（date_key）—— 与 PDF/交互层同一口径

# ---- 几个可复用的样式常量（全文件共用，避免每个单元格重复配置）----
# 四周细边框（Side("thin") 一条边，Border 把四边拼成一个“边框样式”）
THIN = Border(top=Side("thin"), bottom=Side("thin"),
              left=Side("thin"), right=Side("thin"))
C = Alignment(horizontal="center", vertical="center")          # 水平垂直都居中
CW = Alignment(horizontal="center", vertical="center", wrap_text=True)  # 居中+自动换行
L = Alignment(vertical="center")            # B列材料名：水平不定，垂直居中
LW = Alignment(vertical="center", wrap_text=True)             # 材料名可换行

# 字体：name=宋体（中文默认），size=字号，bold=加粗。
F_TITLE = Font(name="宋体", size=24, bold=True)    # 大标题
F_HEAD = Font(name="宋体", size=12)                # 表头
F_MAIN = Font(name="宋体", size=14, bold=True)     # 大类行
F_SUB = Font(name="宋体", size=12, bold=True)      # 小类行
F_BODY = Font(name="宋体", size=12)                # 数据正文

# 模板骨架定义：(行类型, A列类号, B列类名)。这是“表格的长什么样”的总清单。
#   M=大类行  S=小类行（四/九下挂小类）  （真正的材料行由数据填充，编号在各类内重排）
SKELETON = [
    ("M", "一", "履历类材料"),
    ("M", "二", "自传和思想类材料"),
    ("M", "三", "考核鉴定类材料"),
    ("M", "四", ""),                          # 四：无大类名称（只写类号四）
    ("S", "四-1", "学历学位类材料"),
    ("S", "四-2", "专业技术职务（职称）类材料"),
    ("S", "四-3", "学术评鉴类材料"),
    ("S", "四-4", "教育培训类材料"),
    ("M", "五", "政审审计和审核类材料"),
    ("M", "六", "党、团类材料"),
    ("M", "七", "表彰奖励类材料"),
    ("M", "八", "违规违纪违法处理处分类材料"),
    ("M", "九", ""),                          # 九：无大类名称
    ("S", "九-1", "工资类材料"),
    ("S", "九-2", "任免类材料"),
    ("S", "九-3", "出国类材料"),
    ("S", "九-4", "录用招工警衔晋级及代表会议类材料"),
    ("M", "十", "其他可供组织参考的材料"),
]

# 各列宽度（字典：列字母 → 宽度）。D/E 没列 → 用默认。
COL_WIDTHS = {"A": 8.5, "B": 38.4, "C": 8.0, "F": 7.6, "G": 7.0, "H": 8.75}
TAIL_BLANK_ROWS = 6          # 模板末尾保留的空框行数（模板风格，手写目录常留几行空表格）


def export(materials: list[dict], person: str, out_dir: str,
           mark_inferred: bool = True) -> str:
    """主入口：导出《<person>-人事档案目录.xlsx》并返回文件路径。

    流程分四段（对应下面几个注释块）：
      ① 建工作簿 + 设列宽
      ② 画表头（两行合并单元格：类号/材料名称/制成时间[年月日]/份数/页数/备注）
      ③ 沿 SKELETON 走一遍：画大类/小类行，并把对应类别的材料行填进去
      ④ 尾部空框行 + 一个隐藏的“溯源表(_trace)”存每行材料来自哪些照片 → 保存
    """
    wb = Workbook()                 # 新建一个 Excel 文件（内存中）
    ws = wb.active                  # 取默认的第一张表（active sheet）
    ws.title = "档案目录"            # 给这张表起名

    # ① 设列宽：遍历 COL_WIDTHS，把每列的宽度设成我们要的（A 列 8.5 字符宽…）
    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w

    # ---------------- ② 表头（与模板逐格一致） ----------------
    ws.merge_cells("A1:H1")         # 把第一行 A..H 合并成一个格子（放总标题）
    t = ws["A1"]
    t.value = "运城市国有企业退休人员社会化管理\n人事档案目录"   # \n 在合并格内靠 wrap 换行
    t.font = F_TITLE
    t.alignment = CW
    ws.row_dimensions[1].height = 78            # 第一行加高，放得下两行大字

    # 表头第二、三行：列出列名。合并：类号/材料名称各跨两行，制成时间跨 3 列再拆年月日
    heads = [("A2:A3", "类号"), ("B2:B3", "材料名称"),
             ("C2:E2", "材料制成时间"), ("F2:F3", "份数"),
             ("G2:G3", "页数"), ("H2:H3", "备注")]
    for rng, text in heads:
        ws.merge_cells(rng)                     # 合并范围（如 A2:A3）
        cell = ws[rng.split(":")[0]]            # 取范围左上角格写入文字
        cell.value = text
        cell.font = F_HEAD
        cell.alignment = C
    # C2:E2 合并成“材料制成时间”后，第三行再写 年/月/日 三个子列
    for col, text in (("C", "年"), ("D", "月"), ("E", "日")):
        cell = ws[f"{col}3"]
        cell.value = text
        cell.font = F_HEAD
        cell.alignment = C
    # 给第 2、3 行所有单元格加细边框 + 设行高
    for r in (2, 3):
        ws.row_dimensions[r].height = 31 if r == 2 else 28
        for c in range(1, 9):                   # 列 1..8 = A..H
            ws.cell(row=r, column=c).border = THIN

    # ---------------- ③ 骨架 + 数据 ----------------
    # 先把材料按“小类”归组：by_cat[小类] = 属于它的材料行列表
    # setdefault(key, [])：这个 key 还没出现过就给个空 list，再 append —— 不用手动判存在
    by_cat: dict[str, list[dict]] = {}
    for m in materials:
        by_cat.setdefault(m.get("category") or "?", []).append(m)

    row = 4                         # 数据从第 4 行开始写（第 1 行标题、2-3 行表头）
    trace = []                      # 溯源表数据：每行材料记 [序号, 小类, 名称, 源照片seq]
    # 沿 SKELETON 一格一格往下走：它是“表格结构”，我们边画骨架边填材料
    for kind, cat_no, cat_name in SKELETON:
        # 定义“写一行骨架”的小函数（每行：A类号+B类名+全行细边框+行高）
        def _write_row(r, a_val=None, b_val=None, a_font=F_BODY,
                       b_font=F_BODY, a_align=C, b_align=L):
            ca = ws.cell(row=r, column=1)       # A 列单元格
            cb = ws.cell(row=r, column=2)       # B 列单元格
            ca.value = a_val
            cb.value = b_val if b_val != "" else None   # 空字符串写成 None(=空)
            ca.font, cb.font = a_font, b_font
            ca.alignment, cb.alignment = a_align, b_align
            for c in range(1, 9):               # 整行 8 列都加细边框
                ws.cell(row=r, column=c).border = THIN
                if c >= 3:                      # C 之后的列统一居中+正文
                    ws.cell(row=r, column=c).alignment = C
                    ws.cell(row=r, column=c).font = F_BODY
            ws.row_dimensions[r].height = 36    # 每行统一高

        if kind == "M":
            # 大类行（如“一 履历类材料”）：大类字体
            _write_row(row, cat_no, cat_name, a_font=F_MAIN, b_font=F_SUB)
            row += 1
            # 四/九 之外的大类（一/二/三/五/六/七/八/十）材料直接挂在大类行下
            if cat_no not in ("四", "九"):
                row = _write_materials(ws, row, by_cat.get(cat_no, []),
                                       trace, mark_inferred)
        else:  # kind == "S" 小类行（如“四-1 学历学位类材料”）
            _write_row(row, cat_no, cat_name, a_font=F_SUB, b_font=F_SUB,
                       a_align=C, b_align=L)
            row += 1
            # 小类行下面填属于它的材料
            row = _write_materials(ws, row, by_cat.get(cat_no, []),
                                   trace, mark_inferred)

    # ---------------- 尾部空框行（模板风格） ----------------
    for _ in range(TAIL_BLANK_ROWS):            # _ 表示“不需要这个循环变量”
        for c in range(1, 9):
            ws.cell(row=row, column=c).border = THIN
            ws.cell(row=row, column=c).font = F_BODY
        ws.row_dimensions[row].height = 36
        row += 1

    os.makedirs(out_dir, exist_ok=True)         # 输出目录不存在就建
    xlsx = os.path.join(out_dir, f"{person}-人事档案目录.xlsx")   # 拼输出文件名

    # ---------------- 隐藏溯源表（给后期对账/审核用，不看也影响不到） ----------------
    tr = wb.create_sheet("_trace")              # 新建一张隐藏表
    tr.append(["材料序号", "小类", "名称", "源照片seq"])   # 表头行
    for r_ in trace:                            # 每行材料一行
        tr.append(r_)
    tr.sheet_state = "hidden"                   # 标记为隐藏

    wb.save(xlsx)                               # 存盘成 .xlsx 文件
    return xlsx                                 # 把路径交给调用方（adapter 存进 station 文件区）


def _write_materials(ws, row: int, mats: list[dict], trace: list,
                     mark_inferred: bool) -> int:
    """在某（小）类下写“材料行”，编号从 1 开始（与模板一致：每类内独立编号）。

    排序规则：材料按制成时间升序（年→月→日）；缺月的按 1 月算、无日期的垫到最后。
    返回“写完之后游标停在哪一行”，好让上一层接着往下写。
    """
    def _date_key(m):
        """给 sorted 用的“排序键”：日期升序 + 并列时按最小照片号兜底。

        ★ 日期怎么比统一走 classes.date_key（目录序的唯一口径，PDF 与交互层共用同一份），
          这里只多做一件事：**并列时给个确定的兜底**。以前 tied 时按"传进来的顺序"
          （来自大模型的输出顺序，不稳定）→ 同一份数据两次导出可能行序不一样。
          照片号是确定的，拿它兜底最省事。
        """
        ms = m.get("raw_seqs") or m.get("members") or [0]
        return (classes.date_key(m.get("date")), min(ms) if ms else 0)

    mats = sorted(mats, key=_date_key)          # 按时间排好
    for no, m in enumerate(mats, start=1):      # no=类内序号(从1)
        name = m.get("title") or ""             # 材料名
        if not name:
            # 没识别出名字：不能留空导出 → 放一个醒目的待命名提示
            name = "【待确认—无表头，请人工命名】"
        elif mark_inferred and m.get("verdict") == "doubt":
            # 存疑材料（模型自己都没把握）也标“待确认”，逼人工看一眼
            name = f"【待确认】{name}"
        d = m.get("date") or {}
        # 把这一行各列的值准备好：1=序号 2=名称 3/4/5=年月日 6=份数 7=页数
        vals = {1: no, 2: name, 3: d.get("y"), 4: d.get("m"),
                5: d.get("d"), 6: m.get("copies", 1),
                7: m.get("pages", 1)}
        for c, v in vals.items():               # 逐个单元格填值 + 样式
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = F_BODY
            cell.border = THIN
            cell.alignment = LW if c == 2 else C    # 名称列可换行，其它居中
        ws.cell(row=row, column=8).border = THIN    # 备注列(第8列)也画个边
        ws.row_dimensions[row].height = 36
        # 记进溯源表：这一行材料对应源卷里的哪些照片（用 members 页序号）
        trace.append([no, m.get("category"), name,
                      ",".join(map(str, m.get("members", [])))])
        row += 1                                # 游标下移一行
    return row
