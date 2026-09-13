"""weekly-report 技能代码 —— agent 型：收集本周要点 → 生成周报(md/docx) 并落盘到文件区。

tools() 返回叶子 Tool（registry 会补成 weekly.*）。全部本地执行、零外部依赖，
数据源=笔记目录(默认 data/weekly-notes，没有则用仓库内示例) + 本仓 git log，
演示“skill 自带代码 + 多工具串联 + 产物落盘”的宿主扩展路径。

新手视角：这是【第 3 种代码位置】的技能 —— 代码放在技能目录自己的 code/ 下，
registry 会在 import 时把这个目录临时加进 sys.path（原理见 registry.py 的 _import_entry）。
真实“写你自己的技能”最常改的就是这个文件：换工具逻辑 = 改这里的函数 + tools()。
"""
from __future__ import annotations

import datetime
import time

from station import config as _cfg     # 拿仓库/数据路径（config 模块）
from station.core.tool import Tool     # 工具“盒子”（见 tool.py）


def _notes(notes_dir) -> list[str]:
    """读一个目录下所有 .md，每个文件抽一句“标题/首行”返回。"""
    from pathlib import Path            # 用到才 import
    d = Path(notes_dir)
    if not d.is_dir():                  # 目录不存在就返回空列表
        return []
    out = []
    # sorted(d.glob("*.md"))：按文件名排序扫所有 md 文件（glob≈通配符匹配）
    for p in sorted(d.glob("*.md")):
        try:
            text = p.read_text(encoding="utf-8")
        except Exception:               # noqa  # 个别文件读不动就跳过，别中断整个任务
            text = ""
        head = "（无标题）"
        # 取第一个非空行当“一句话要点”（去掉开头的 #）
        for ln in text.splitlines():
            s = ln.strip()
            if s:
                head = s.lstrip("#").strip()[:80]   # lstrip 去掉前导的 #；截 80 字
                break
        out.append(f"- {p.stem}：{head}")           # p.stem = 不含后缀的文件名
    return out


def _git_log(days: int) -> list[str]:
    """跑 git 拿最近 N 天提交。用 subprocess≈Runtime.exec：起一个外部命令并读它的输出。"""
    since = (datetime.date.today() - datetime.timedelta(days=max(1, days)))
    try:
        import subprocess
        # 命令拆成一个 list（不拼字符串）→ 避免空格/注入问题，这是推荐写法
        r = subprocess.run(
            ["git", "-C", str(_cfg.REPO), "log", f"--since={since.isoformat()}",
             "--pretty=format:%ad %h %s", "--date=format:%m-%d"],
            capture_output=True, text=True, encoding="utf-8", timeout=10)
        # splitlines 按行拆；过滤空行
        return [ln for ln in r.stdout.splitlines() if ln.strip()]
    except Exception:                   # noqa  # 没装 git / 不是 git 仓库都别崩
        return []


def _owner(ctx) -> str:
    """产物归属：对话场景用户 id 在 Thread 上；job/脚本场景 Context 直接带。"""
    t = getattr(ctx, "thread", None)
    return (getattr(t, "user_id", "") if t else "") or getattr(ctx, "user_id", "") or ""


def _week_from_md(md: str) -> str:
    """从周报模板首行《# 工作周报 · YYYY-Www》里反推周标签，兜底取本周。"""
    for ln in (md or "").splitlines():
        s = ln.strip()
        if s.startswith("# 工作周报 · "):
            return s.split("·", 1)[1].strip()
    return time.strftime("%Y-W%W")


def collect_highlights(ctx, days: int = 7, notes_dir: str = "") -> str:
    """工具①：收集本周素材（笔记 + git 提交），返回一段 markdown 给模型当“材料”。"""
    if not notes_dir:                   # 没指定笔记目录 → 自动选：用户 data 优先，否则示例
        cand = _cfg.DATA_DIR / "weekly-notes"
        cand2 = _cfg.REPO / "skills" / "weekly-report" / "notes"
        notes_dir = cand if cand.is_dir() else cand2
    notes = _notes(notes_dir)
    git = _git_log(int(days or 7))
    # 组一段结构化 markdown 返回（agent 拿到后决定怎么写进周报）
    parts = [f"## 本周记录（近 {int(days)} 天）", "### 笔记要点"]
    parts.extend(notes or ["（暂无笔记，可先建 data/weekly-notes/*.md）"])
    parts.append("### git 提交")
    parts.extend(git or ["（暂无 git 提交）"])
    return "\n".join(parts)


def render_report(ctx, content: str = "", week: str = "") -> str:
    """工具②：把素材渲染成《工作周报》md 并落盘到文件区，返回下载信息。

    落盘用 station 的 files/store（≈上传到“文件柜”，拿到 id 前端就能下载）。
    """
    from station.files import store as fs
    week = week.strip() or time.strftime("%Y-W%W")   # 空就取“今年-第几周”
    body = (content or "").strip()
    # 让内容读起来像要点列表：如果是一整句话，就包成一条“- xxx”
    if body and not body.startswith("-") and "\n" not in body:
        body = f"- {body}"
    # 组周报 markdown（模板固定骨架）
    md = "\n".join([
        f"# 工作周报 · {week}", "",
        "## 本周完成", body or "（本周没有新增内容）", "",
        "## 数据 / 进展", "（可补充量化数据）", "",
        "## 遇到的问题", "（补充）", "",
        "## 下周计划", "- [ ] 待定", ""])
    fid = fs.save_text(md, ".md", name=f"周报-{week}.md",
                       owner=_owner(ctx), skill_id="weekly-report",
                       group_key=f"weekly:{week}",
                       group_label=f"写周报 · {week}")       # 存进文件区拿 id
    return f"周报已生成并落盘：周报-{week}.md（id={fid}）→ 下载 GET /api/files/{fid}"


def export_docx(ctx, file_id: str = "", content: str = "") -> str:
    """工具③（可选）：把 md 周报转 .docx。python-docx 没装就给提示，别崩。"""
    from station.files import store as fs
    try:
        from docx import Document      # 第三方库；装没装用 try 探一下
    except Exception:                  # noqa
        return "docx 导出不可用：python-docx 未安装（直接用 .md 产物即可）"
    # 输入二选一：给 file_id → 读那个文件；否则直接用 content
    if file_id:
        p = fs.path(file_id)
        md = p.read_text(encoding="utf-8") if p else ""
    else:
        md = content or ""
    doc = Document()                   # 建 Word 文档对象
    for ln in md.splitlines():         # 逐行解析 markdown → Word 段落/标题/列表
        if ln.startswith("# "):        # 一级标题 → Word Heading 1
            doc.add_heading(ln[2:].strip(), level=1)
        elif ln.startswith("## "):
            doc.add_heading(ln[3:].strip(), level=2)
        elif ln.strip() == "":         # 空行忽略
            pass
        elif ln.strip().startswith("- "):        # 列表项 → Word 项目符号
            doc.add_paragraph(ln.strip()[2:], style="List Bullet")
        else:
            doc.add_paragraph(ln.strip())        # 普通行 → 正文
    import io
    buf = io.BytesIO()                 # 先在内存里存成字节（不急着写盘）
    doc.save(buf)
    week = _week_from_md(md)
    name = f"周报-{week}.docx"
    fid = fs.save_bytes(buf.getvalue(), ".docx", name=name,
                        owner=_owner(ctx), skill_id="weekly-report",
                        group_key=f"weekly:{week}",
                        group_label=f"写周报 · {week}")      # 字节存进文件区
    return f"已导出 docx：{name}（id={fid}）→ 下载 GET /api/files/{fid}"


def tools() -> list[Tool]:
    """暴露本技能的工具清单（registry 会注册成 weekly.*）。写周报的“三步” = 这三个工具。"""
    return [
        Tool(name="collect_highlights", label="收集本周素材",
             description="收集本周素材：扫笔记目录(data/weekly-notes 或示例) + 本仓 git 近 N 天提交，"
                         "返回待整理的要点 markdown。写周报前务必先调它拿素材。",
             run=collect_highlights,
             args=[{"name": "days", "type": "int", "desc": "回溯天数(默认7)", "required": False},
                   {"name": "notes_dir", "type": "str", "desc": "笔记目录(可省)", "required": False}]),
        Tool(name="render_report", label="生成周报",
             description="把素材内容渲染成一份《工作周报》markdown 并落盘到文件区，返回产物下载信息。",
             run=render_report,
             args=[{"name": "content", "type": "str", "desc": "周报正文要点", "required": True},
                   {"name": "week", "type": "str", "desc": "周标签(可省，默认本周)", "required": False}]),
        Tool(name="export_docx", label="导出 Word",
             description="把已生成的 md 周报(file_id)转成 .docx 产物（需 python-docx，可选）。",
             run=export_docx,
             args=[{"name": "file_id", "type": "str", "desc": "md 产物 id", "required": False},
                   {"name": "content", "type": "str", "desc": "或直接给正文", "required": False}]),
    ]
